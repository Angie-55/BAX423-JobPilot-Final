"""DOCX export helpers for JobPilot resume drafts."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


FORBIDDEN_RESUME_TEXT = [
    "Tailoring Notes",
    "raw JSON",
    "debug",
    "[insert",
    "Company Name",
    "add metrics here",
    "Use the original resume facts",
    "TBD",
    "placeholder",
]


def clean_resume_markdown(markdown_text: str) -> str:
    text = str(markdown_text or "").replace("\r\n", "\n").strip()
    for forbidden in FORBIDDEN_RESUME_TEXT:
        text = text.replace(forbidden, "")
    return text


def configure_resume_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size in [("Heading 1", 12), ("Heading 2", 11)]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(3)


def add_resume_paragraph(doc: Document, line: str, is_contact_line: bool = False) -> None:
    stripped = line.strip()
    if not stripped:
        return
    if stripped.startswith("- "):
        paragraph = doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        return
    paragraph = doc.add_paragraph(stripped)
    if is_contact_line:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(8)


def markdown_resume_to_docx(markdown_text: str) -> Document:
    doc = Document()
    configure_resume_doc(doc)
    text = clean_resume_markdown(markdown_text)
    lines = [line.rstrip() for line in text.splitlines()]
    seen_title = False
    previous_was_title = False

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(18)
            paragraph.paragraph_format.space_after = Pt(2)
            seen_title = True
            previous_was_title = True
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            previous_was_title = False
            continue
        add_resume_paragraph(doc, line, is_contact_line=seen_title and previous_was_title)
        previous_was_title = False

    return doc


def save_resume_docx(markdown_text: str, output_path: str | Path) -> None:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = markdown_resume_to_docx(markdown_text)
    doc.save(path)


def resume_docx_bytes(markdown_text: str) -> bytes:
    buffer = BytesIO()
    doc = markdown_resume_to_docx(markdown_text)
    doc.save(buffer)
    return buffer.getvalue()
